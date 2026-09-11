"""
docx_io.py (originally utils.py in the standalone TOC Formatter Tool repo,
renamed here only to avoid colliding with this project's own utils/ package)
--------
DOCX I/O helpers for the TOC Formatter Tool. Reading logic lives here so
formatter.py stays focused purely on numbering/formatting rules (the part
most likely to change in future). The UI layer should not touch python-docx
directly - it always goes through this module.
"""

import io
import os
import re
import zipfile
from docx import Document
from docx.oxml.ns import qn

from toc_tool import formatter

# Word paragraph styles that carry an explicit outline depth in their name,
# e.g. "TOC 1", "TOC 2", "Heading 1", "heading3". This is the MOST reliable
# hierarchy signal for a real Word document's Table of Contents, because a
# huge number of real-world TOCs have NO literal "1.1." numbers typed into
# the text at all - the numbering (if any) is rendered by Word's own
# multilevel-list engine and is never part of para.text. Without reading
# this, every line looks "unnumbered" and the tool has no choice but to
# flatten everything to one level (which is the bug this fixes).
_STYLE_LEVEL_RE = re.compile(r"(?:toc|heading)\s*0*([1-9])", re.IGNORECASE)


def _style_level(style_name: str):
    if not style_name:
        return None
    m = _STYLE_LEVEL_RE.search(style_name)
    return int(m.group(1)) if m else None


def _outline_level(paragraph):
    """Read w:pPr/w:outlineLvl (0-based) directly from the paragraph XML, if set."""
    pPr = paragraph._p.pPr
    if pPr is None:
        return None
    el = pPr.find(qn("w:outlineLvl"))
    if el is None:
        return None
    val = el.get(qn("w:val"))
    try:
        return int(val)
    except (TypeError, ValueError):
        return None


def _list_ilvl(paragraph):
    """Read w:pPr/w:numPr/w:ilvl (0-based) directly from the paragraph XML, if set."""
    pPr = paragraph._p.pPr
    if pPr is None:
        return None
    numPr = pPr.find(qn("w:numPr"))
    if numPr is None:
        return None
    ilvl_el = numPr.find(qn("w:ilvl"))
    if ilvl_el is None:
        return None
    val = ilvl_el.get(qn("w:val"))
    try:
        return int(val)
    except (TypeError, ValueError):
        return None


def _structural_level(paragraph) -> int:
    """
    Best-effort structural nesting level for a paragraph, checked in priority
    order: named style (TOC n / Heading n) > explicit outline level > list
    level (numPr/ilvl). Returns None if none of these are present, in which
    case formatter.py falls back to parsing the visible text itself.
    """
    style_name = paragraph.style.name if paragraph.style is not None else ""
    lvl = _style_level(style_name)
    if lvl is not None:
        return lvl

    lvl = _outline_level(paragraph)
    if lvl is not None:
        return lvl + 1  # outlineLvl is 0-based

    lvl = _list_ilvl(paragraph)
    if lvl is not None:
        return lvl + 1  # ilvl is 0-based

    return None


# ---------------------------------------------------------------------------
# READING
# ---------------------------------------------------------------------------

def extract_lines_from_docx(file_like) -> list:
    """
    Read a .docx file (path str, or file-like/BytesIO object) and return a
    list of (text, structural_level_or_None) tuples, in document order.

    structural_level comes from Word's own formatting (Heading/TOC style,
    outline level, or list level) when available - this is the ground truth
    for real Word documents. formatter.parse_toc() falls back to parsing the
    visible text (numbers like "1.1." or legacy "A./a./i./1." markers) only
    when no structural level is present, e.g. for plain pasted text.
    """
    doc = Document(file_like)
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append((text, _structural_level(para)))

    # Also walk tables, in case the TOC was pasted into a table (common in
    # some report templates) - append after body paragraphs.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    text = para.text.strip()
                    if text:
                        lines.append((text, _structural_level(para)))
    return lines


def docx_to_result(file_like, fmt: str = "txt"):
    """
    Full single-file pipeline: docx (path or file-like) -> raw (text, level)
    pairs -> formatter.process_text -> (entries, rendered_string).
    """
    lines = extract_lines_from_docx(file_like)
    entries, rendered = formatter.process_text(lines, fmt=fmt)
    return entries, rendered


# ---------------------------------------------------------------------------
# BATCH PROCESSING (Phase 2)
# ---------------------------------------------------------------------------

def batch_process(uploaded_files, fmt: str = "txt", progress_callback=None):
    """
    Process many uploaded docx files (Streamlit UploadedFile objects, or
    (name, bytes)/(name, file-like) tuples).

    progress_callback(done_count, total_count, current_filename) is called
    after each file, so the UI can drive a progress bar.

    Returns a list of dicts: {"filename": ..., "output_name": ...,
                               "rendered": ..., "error": str|None}
    """
    results = []
    total = len(uploaded_files)

    for i, uf in enumerate(uploaded_files, start=1):
        name = getattr(uf, "name", None) or (uf[0] if isinstance(uf, tuple) else f"file_{i}.docx")
        source = uf if not isinstance(uf, tuple) else uf[1]

        base = os.path.splitext(os.path.basename(name))[0]
        ext = {"txt": "txt", "md": "md", "json": "json"}.get(fmt, "txt")
        output_name = f"{base}_formatted.{ext}"

        try:
            # Streamlit UploadedFile / file-like objects need seek(0) if reused
            if hasattr(source, "seek"):
                source.seek(0)
            _, rendered = docx_to_result(source, fmt=fmt)
            results.append({
                "filename": name,
                "output_name": output_name,
                "rendered": rendered,
                "error": None,
            })
        except Exception as exc:  # noqa: BLE001 - surface per-file errors, keep batch going
            results.append({
                "filename": name,
                "output_name": output_name,
                "rendered": "",
                "error": str(exc),
            })

        if progress_callback:
            progress_callback(i, total, name)

    return results


def results_to_zip_bytes(results) -> bytes:
    """Package batch_process() results into an in-memory ZIP file's bytes."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for r in results:
            if r["error"]:
                continue
            zf.writestr(r["output_name"], r["rendered"])
    buf.seek(0)
    return buf.getvalue()
